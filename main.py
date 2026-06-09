# main.py
import csv
from io import StringIO
from datetime import datetime
import pandas as pd

from db.init_db import init_db
from db.session import SessionLocal
from models.transaction import Transaction
from models.merchant import Merchant
from importers.names_accounts import import_names
from cleaner.normalizer import normalize
from cleaner.matcher import match_transaction
from config import MATCH_BY
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


def load_transactions(filepath: str):
    """Read a Chase CSV and insert raw transactions into the database."""
    print(f"\nReading {filepath}...")

    with open(filepath, encoding='latin-1') as f:
        content = f.read().replace('\x00', '')

    reader = csv.DictReader(StringIO(content))
    rows = list(reader)

    session = SessionLocal()

    # Wipe existing transactions so the DB reflects only the file just loaded.
    deleted = session.query(Transaction).delete()
    if deleted:
        print(f"Cleared {deleted} existing transactions.")

    imported = 0
    skipped = 0

    for row in rows:
        description = str(row.get("Description", "")).strip()
        amount = row.get("Amount", "").strip()
        date_raw = str(row.get("Posting Date", "")).strip()
        transaction_type = str(row.get("Type", "")).strip()

        if not description:
            skipped += 1
            continue

        try:
            amount = float(amount)
        except (ValueError, TypeError):
            skipped += 1
            continue

        try:
            date = datetime.strptime(date_raw, "%m/%d/%Y").date()
        except ValueError:
            skipped += 1
            continue

        transaction = Transaction(
            date=date,
            raw_description=description,
            amount=amount,
            source="Chase",
            transaction_type=transaction_type,
        )
        session.add(transaction)
        imported += 1

    session.commit()
    session.close()
    print(f"Done. {imported} imported, {skipped} skipped.")


def clean_transactions():
    """Normalize and match all uncleaned transactions."""
    session = SessionLocal()
    unclean = session.query(Transaction).filter_by(is_cleaned=False).all()

    if not unclean:
        print("\nNo uncleaned transactions found.")
        session.close()
        return

    print(f"\nCleaning {len(unclean)} transactions...")
    matched = 0
    unmatched = 0

    for transaction in unclean:
        transaction.clean_description = normalize(
            transaction.raw_description,
            transaction.transaction_type
        )
        transaction.is_cleaned = True

        found = match_transaction(transaction)
        if found:
            matched += 1
        else:
            unmatched += 1

    session.commit()
    session.close()
    print(f"Done. {matched} matched, {unmatched} unmatched.")


def review_unmatched():
    """Print all transactions that could not be matched to a merchant."""
    session = SessionLocal()
    unmatched = session.query(Transaction).filter_by(is_matched=False).all()

    if not unmatched:
        print("\nNo unmatched transactions.")
        session.close()
        return

    print(f"\n{'Date':<12} {'Amount':>10}  {'Clean Description':<40} {'Raw Description'}")
    print("-" * 100)
    for t in unmatched:
        print(f"{str(t.date):<12} {t.amount:>10.2f}  {str(t.clean_description):<40} {t.raw_description}")

    session.close()


def export_to_excel():
    """Export all transactions to an Excel file."""
    session = SessionLocal()
    transactions = session.query(Transaction).order_by(Transaction.date).all()
    session.close()

    if not transactions:
        print("\nNo transactions to export.")
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "Transactions"

    # Header row
    headers = ["Date", "Raw Description", "Clean Description", "Account Number", "Account Name", "Amount", "Matched"]
    ws.append(headers)

    # Style the header
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF", name="Arial")
        cell.fill = PatternFill("solid", start_color="2F4F8F")
        cell.alignment = Alignment(horizontal="center")

    # Data rows
    for t in transactions:
        ws.append([
            str(t.date),
            t.raw_description,
            t.clean_description or "",
            t.account_number or "",
            t.account_name or "",
            t.amount,
            "Yes" if t.is_matched else "No",
        ])

    # Highlight unmatched rows in light yellow
    yellow = PatternFill("solid", start_color="FFFACD")
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        if row[6].value == "No":
            for cell in row:
                cell.fill = yellow

    # Column widths
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 50
    ws.column_dimensions["C"].width = 35
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["E"].width = 28
    ws.column_dimensions["F"].width = 12
    ws.column_dimensions["G"].width = 10

    # Amount column — currency format
    for row in ws.iter_rows(min_row=2, min_col=6, max_col=6):
        for cell in row:
            cell.number_format = '#,##0.00;(#,##0.00)'

    output_path = "data/cleaned_transactions.xlsx"
    wb.save(output_path)
    print(f"\nExported {len(transactions)} transactions to {output_path}")

def setup():
    """Make sure the database exists and merchants are seeded."""
    init_db()
    session = SessionLocal()
    has_merchants = session.query(Merchant).first() is not None
    session.close()
    if not has_merchants:
        import_names()


def main():
    print("\n" + "=" * 40)
    print("       FEEDS CLEANER")
    print("=" * 40)

    setup()

    raw = input("\nEnter path to CSV file (or 'q' to quit): ").strip()
    if raw.lower() in ("", "q", "quit", "exit"):
        print("\nGoodbye.\n")
        return

    # Tolerate paths pasted with surrounding quotes.
    filepath = raw.strip('"').strip("'")

    load_transactions(filepath)
    clean_transactions()
    export_to_excel()


if __name__ == "__main__":
    main()